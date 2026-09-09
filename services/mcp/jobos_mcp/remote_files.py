"""Bounded, content-addressed files inside the existing publication inbox."""

from __future__ import annotations

import base64
import io
import os
import re
import zipfile
from contextlib import suppress
from hashlib import sha256
from pathlib import Path
from xml.sax.saxutils import escape

from jobos_mcp.server import (
    _prepare_document_publication_workspace,
    _read_publication_input,
)

MAX_TRANSFER = 2_000_000
FILE_ID = re.compile(r"^[a-f0-9]{64}\.(?:txt|md|json|pdf|docx)$")


def put_file(root: Path, conversation_id: str, job_id: str, suffix: str, data: bytes) -> dict:
    if suffix not in {"txt", "md", "json", "pdf", "docx"}:
        raise ValueError("Supported formats: txt, md, json, pdf, docx")
    if not 0 < len(data) <= MAX_TRANSFER:
        raise ValueError("File must contain 1 to 2000000 bytes")
    file_id = f"{sha256(data).hexdigest()}.{suffix}"
    workspace = _prepare_document_publication_workspace(conversation_id, job_id, artifact_root=root)
    # Anchor every component at the artifact root; never follow a swapped inbox.
    descriptors = []
    flags = os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW
    temporary = f".upload-{os.urandom(16).hex()}"
    try:
        current = os.open(root, flags)
        descriptors.append(current)
        for part in workspace.relative_to(root.resolve()).parts:
            current = os.open(part, flags, dir_fd=current)
            descriptors.append(current)
        fd = os.open(
            temporary, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW, 0o600, dir_fd=current
        )
        try:
            with os.fdopen(fd, "wb") as stream:
                stream.write(data)
                stream.flush()
                os.fsync(stream.fileno())
            # Existing immutable files are verified by the read below.
            with suppress(FileExistsError):
                os.link(
                    temporary,
                    file_id,
                    src_dir_fd=current,
                    dst_dir_fd=current,
                    follow_symlinks=False,
                )
        finally:
            os.unlink(temporary, dir_fd=current)
        os.fsync(current)
    finally:
        for fd in reversed(descriptors):
            os.close(fd)
    read_file(root, conversation_id, job_id, file_id)
    return {"file_id": file_id, "sha256": sha256(data).hexdigest(), "size_bytes": len(data)}


def read_file(root: Path, conversation_id: str, job_id: str, file_id: str) -> bytes:
    if not FILE_ID.fullmatch(file_id):
        raise ValueError("Invalid file ID; use the ID returned by file_upload or document_generate")
    workspace = _prepare_document_publication_workspace(conversation_id, job_id, artifact_root=root)
    _, data = _read_publication_input(
        str(workspace / file_id),
        conversation_id=conversation_id,
        job_id=job_id,
        artifact_root=root,
        maximum=MAX_TRANSFER,
    )
    if sha256(data).hexdigest() != file_id.split(".")[0]:
        raise ValueError("File checksum changed")
    return data


def decode_upload(content: str, encoding: str) -> bytes:
    if encoding == "text":
        data = content.encode("utf-8")
    elif encoding == "base64":
        if len(content) > ((MAX_TRANSFER + 2) // 3) * 4:
            raise ValueError("File exceeds 2000000 bytes")
        try:
            data = base64.b64decode(content, validate=True)
        except (ValueError, UnicodeError) as error:
            raise ValueError("Invalid base64 content") from error
    else:
        raise ValueError("Encoding must be text or base64")
    if not 0 < len(data) <= MAX_TRANSFER:
        raise ValueError("File must contain 1 to 2000000 bytes")
    return data


def generate_pair(markdown: str) -> tuple[bytes, bytes]:
    """Simple text, # headings and - bullets, not HTML or a browser renderer."""
    import reportlab
    from docx import Document
    from docx.shared import Inches, Pt
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    if not markdown.strip() or len(markdown.encode("utf-8")) > 100_000:
        raise ValueError("Markdown must contain 1 to 100000 UTF-8 bytes")
    font = TTFont("JobOSVera", str(Path(reportlab.__file__).parent / "fonts" / "Vera.ttf"))
    # Locate the font shipped by ReportLab, independent of host font installations.
    pdfmetrics.registerFont(font)
    if any(ord(c) not in font.face.charToGlyph for c in markdown if not c.isspace()):
        raise ValueError("Text contains characters unsupported by the bundled PDF font")
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.65)
    doc.styles["Normal"].font.size = Pt(10)
    story = []
    normal = ParagraphStyle(
        "JobOS", fontName="JobOSVera", fontSize=10, leading=14, spaceAfter=6, splitLongWords=True
    )
    heading = ParagraphStyle("JobOSHeading", parent=normal, fontSize=14, leading=18, spaceBefore=8)
    for line in markdown.splitlines():
        if not line.strip():
            story.append(Spacer(1, 5))
            continue
        match = re.match(r"^(#{1,3})\s+(.*)$", line)
        bullet = line.startswith(("- ", "* "))
        text = match[2] if match else line[2:] if bullet else line
        if match:
            doc.add_heading(text, level=len(match[1]))
        else:
            doc.add_paragraph(text, style="List Bullet" if bullet else None)
        story.append(
            Paragraph(escape(("• " if bullet else "") + text), heading if match else normal)
        )
    pdf = io.BytesIO()
    SimpleDocTemplate(
        pdf,
        pagesize=(612, 792),
        topMargin=47,
        bottomMargin=47,
        leftMargin=54,
        rightMargin=54,
        invariant=1,
    ).build(story)
    raw_docx = io.BytesIO()
    doc.save(raw_docx)
    # Stable ZIP timestamps make retried generation have the same content identities.
    docx = io.BytesIO()
    with zipfile.ZipFile(raw_docx) as source, zipfile.ZipFile(docx, "w") as target:
        for item in source.infolist():
            item.date_time = (2000, 1, 1, 0, 0, 0)
            target.writestr(item, source.read(item.filename))
    return pdf.getvalue(), docx.getvalue()
